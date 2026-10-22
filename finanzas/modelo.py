#!/usr/bin/env python3
"""
Modelo financiero de la agencia — cálculos reproducibles.

Uso:
  python3 finanzas/modelo.py                 # informe completo de cada idea en IDEAS
  python3 finanzas/modelo.py --comparar      # tabla comparativa (mismo método para todas)
  python3 finanzas/modelo.py --xlsx          # regenera finanzas/modelo.xlsx con fórmulas
  python3 finanzas/modelo.py --json          # resultados en JSON (para otros informes)
  python3 finanzas/modelo.py --idea "Nombre" # solo una idea
  python3 finanzas/modelo.py --desempate            # Fase 4: margen real por hora (MRH) de A y B con DATOS_DESEMPATE
  python3 finanzas/modelo.py --desempate --ejemplo  # mismo cálculo con las cifras ILUSTRATIVAS del plan (no son datos)
  python3 finanzas/modelo.py --desempate --dato A.adelantos_cobrados=100 --dato B.paso_B2=si ...  # datos por argumento

Reglas (ver CLAUDE.md y .claude/agents/analista-financiero.md):
  * Nadie hace aritmética "de memoria": toda cifra de un informe sale de este script.
  * Cada entrada lleva `origen`: "VERIFICADO: fuente, fecha" | "ESTIMADO: método"
    | "SUPUESTO: por validar" | "DESCONOCIDO". El valor "EJEMPLO" solo es válido
    mientras ES_EJEMPLO = True.
  * Las ideas de ejemplo se borraron el 2026-10-21 (Fase 3 de /explorar). Las cinco ideas
    cargadas abajo son las de la lista corta confirmada en expediente/decisiones.md.
    El script se niega a correr con ES_EJEMPLO = False si queda alguna entrada con origen EJEMPLO.
  * Las proyecciones son escenarios, no predicciones. La sección de sensibilidad dice
    qué entrada mueve más el resultado.
  * Estructura idéntica para las cinco ideas (mismo método, Fase 3):
      - precio con IGV incluido; régimen tributario uniforme SUPUESTO (RER: IGV 18 % + renta 1,5 %)
        hasta que exista informe legal; variantes Nuevo RUS y 4.ª categoría se muestran aparte;
      - pasarela 0 % (Yape/Plin/transferencia) como base; variante Culqi aparte;
      - horas/semana = horas fijas + horas por unidad × unidades del mes 12 del escenario base;
      - chequeo contra F1 (inversión ≤ S/ 400, fijos ≤ S/ 150), F2 (≤ 34 h/semana) y
        capital real (S/ 400 + S/ 350/mes);
      - unidades/mes y mes en que el flujo neto mensual cubre un sueldo (S/ 2 000, supuesto).
Montos en soles (S/). Meses: 12.
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from dataclasses import dataclass, field, asdict

MESES = 12
SEMANAS_POR_MES = 52 / 12
ETIQUETAS_VALIDAS = ("VERIFICADO", "ESTIMADO", "SUPUESTO", "DESCONOCIDO")

# =============================================================================
# 1. ESTRUCTURAS
# =============================================================================


@dataclass
class Entrada:
    """Una cifra de entrada con su etiqueta de origen."""

    valor: float
    origen: str
    nota: str = ""

    def etiqueta(self) -> str:
        return self.origen.split(":")[0].strip().upper()


@dataclass
class Escenario:
    nombre: str
    unidades_mes1: Entrada  # unidades vendidas en el mes 1
    crecimiento_mensual: Entrada  # fracción, ej. 0.10 = +10 % por mes (neto de bajas)
    factor_precio: Entrada  # 1.0 = precio de lista; 0.9 = 10 % de descuento efectivo


def unidades_en_mes(esc: Escenario, m: int) -> float:
    return esc.unidades_mes1.valor * (1 + esc.crecimiento_mensual.valor) ** (m - 1)


@dataclass
class Idea:
    nombre: str
    unidad: str  # qué es "una unidad": una mensualidad, un servicio, un asiento, un pedido, un evento
    precio_venta: Entrada  # precio al cliente, con IGV incluido si igv_aplica
    igv_aplica: bool
    tasa_igv: Entrada
    tasa_pasarela: Entrada  # comisión % de la pasarela sobre el monto cobrado
    fijo_pasarela: Entrada  # S/ fijo por transacción
    tasa_impuesto_renta: Entrada  # % sobre ingreso neto (según régimen que indique legal)
    inversion_inicial: dict[str, Entrada]
    costos_fijos_mensuales: dict[str, Entrada]
    costo_variable_unitario: dict[str, Entrada]  # insumos, envío, IA/tokens, servidor por unidad
    escenarios: dict[str, Escenario]
    horas_fijas_semana: Entrada  # ventas, contenido, administración, mantenimiento
    horas_por_unidad: Entrada  # horas-persona por unidad vendida
    rango_precio: tuple[float, float] | None = None  # (mín, máx) del precio cuando el precio es supuesto
    horas_semana_requeridas: Entrada | None = None  # se calcula en __post_init__ si no se da
    notas: str = ""

    def __post_init__(self):
        if self.horas_semana_requeridas is None:
            u12 = unidades_en_mes(self.escenarios["base"], MESES)
            self.horas_semana_requeridas = Entrada(
                self.horas_fijas_semana.valor + self.horas_por_unidad.valor * u12 / SEMANAS_POR_MES,
                "ESTIMADO: horas fijas/semana + horas por unidad × unidades del mes 12 del escenario base "
                "÷ 4,33 semanas/mes (pico del año 1)")

    def horas_semana_en_mes(self, esc: Escenario, m: int) -> float:
        return self.horas_fijas_semana.valor + self.horas_por_unidad.valor * unidades_en_mes(esc, m) / SEMANAS_POR_MES

    def entradas(self):
        """Itera (ruta, Entrada) para validación."""
        for k in ("precio_venta", "tasa_igv", "tasa_pasarela", "fijo_pasarela",
                  "tasa_impuesto_renta", "horas_fijas_semana", "horas_por_unidad", "horas_semana_requeridas"):
            yield k, getattr(self, k)
        for grupo in ("inversion_inicial", "costos_fijos_mensuales", "costo_variable_unitario"):
            for k, e in getattr(self, grupo).items():
                yield f"{grupo}.{k}", e
        for nombre, esc in self.escenarios.items():
            yield f"escenarios.{nombre}.unidades_mes1", esc.unidades_mes1
            yield f"escenarios.{nombre}.crecimiento_mensual", esc.crecimiento_mensual
            yield f"escenarios.{nombre}.factor_precio", esc.factor_precio


def E(valor: float, origen: str = "EJEMPLO", nota: str = "") -> Entrada:
    return Entrada(valor, origen, nota)


def escenarios_estandar(pes: tuple, base: tuple, opt: tuple, origen) -> dict[str, Escenario]:
    """Cada tupla: (unidades_mes1, crecimiento_mensual, factor_precio).
    `origen`: un texto para los tres escenarios o un dict {"pesimista": ..., "base": ..., "optimista": ...}."""
    out = {}
    for nombre, (u, g, f) in (("pesimista", pes), ("base", base), ("optimista", opt)):
        o = origen[nombre] if isinstance(origen, dict) else origen
        out[nombre] = Escenario(nombre, E(u, o), E(g, o), E(f, o))
    return out


# =============================================================================
# 2. ENTRADAS — Fase 3 de /explorar, 2026-10-21. Cinco ideas de la lista corta.
# =============================================================================
ES_EJEMPLO = False  # las ideas de ejemplo fueron borradas el 2026-10-21

# --- Parámetros globales (mismos para las cinco ideas) -------------------------
CAPITAL_INICIAL = E(400, "VERIFICADO: expediente/00-perfil-socios.md (Indira S/ 400; Carlos S/ 0), 2026-10-21")
APORTE_MENSUAL = E(350, "VERIFICADO: expediente/00-perfil-socios.md (Carlos 150 + Indira 200), 2026-10-21")
TOPE_INVERSION_F1 = E(400, "VERIFICADO: expediente/decisiones.md, filtro F1, 2026-10-21")
TOPE_FIJOS_F1 = E(150, "VERIFICADO: expediente/decisiones.md, filtro F1 (gasto mensual hasta la 1.ª venta), 2026-10-21")
TOPE_HORAS_F2 = E(34, "VERIFICADO: expediente/decisiones.md, filtro F2 (34 h/semana entre los dos en exámenes), 2026-10-21")
TIPO_CAMBIO = E(3.44, "VERIFICADO: open.er-api.com, 2026-10-21 (3,4376 S/ por US$, redondeado)")
RMV = E(1130, "VERIFICADO: D.S. 006-2024-TR, RMV S/ 1 130 vigente desde 2025-01-01; no se verificó hoy si hubo aumento en 2026")
SUELDO_OBJETIVO = E(2000, "SUPUESTO: por validar — sueldo neto mensual de UN socio; los sueldos reales de Carlos e Indira no están "
                          "registrados [DESCONOCIDO]. Se usa S/ 2 000 (entre la RMV de S/ 1 130 y el promedio de community manager "
                          "S/ 1 346 citado en el informe del explorador, con margen para un puesto con estudios). Los socios deben confirmarlo")
CUOTA_RUS_CAT1 = E(20, "VERIFICADO: D. Leg. 937 modificado por D. Leg. 1270, Nuevo RUS categoría 1 (ingresos ≤ S/ 5 000/mes), "
                       "cuota S/ 20; norma conocida, no verificada hoy — confirmar con asesor-legal-tributario-peru")
TASA_RENTA_4TA = E(0.08, "VERIFICADO: retención de rentas de 4.ª categoría 8 % (recibos por honorarios > S/ 1 500); "
                         "norma conocida, no verificada hoy — confirmar con asesor legal; solo aplica a servicios personales (ideas 2 y 3)")
PASARELA_ALT_TASA = E(0.0344, "VERIFICADO: culqi.com/precios, 2026-10-21 — online y link de pago, tarjetas nacionales y billeteras")
PASARELA_ALT_FIJO = E(0.20 * TIPO_CAMBIO.valor, "VERIFICADO: culqi.com/precios, 2026-10-21 — US$ 0,20 por transacción × tipo de cambio")


def tributos_uniformes() -> dict:
    """Mismo bloque tributario y de cobro para las cinco ideas. Sin informe legal todavía."""
    return dict(
        igv_aplica=True,
        tasa_igv=E(0.18, "VERIFICADO: TUO Ley del IGV (16 %) + IPM (2 %) = 18 %, vigente; sin informe legal aún — "
                         "confirmar con asesor-legal-tributario-peru", "precio al cliente con IGV incluido"),
        tasa_pasarela=E(0.0, "SUPUESTO: por validar — cobro por Yape/Plin o transferencia entre personas naturales, sin comisión; "
                             "la variante con Culqi (3,44 % + US$ 0,20) se calcula aparte"),
        fijo_pasarela=E(0.0, "SUPUESTO: por validar — igual que tasa_pasarela"),
        tasa_impuesto_renta=E(0.015, "SUPUESTO: por validar — RER (1,5 % de ingresos netos mensuales) como régimen uniforme y "
                                     "conservador para las cinco ideas; permite emitir factura a negocios. Variantes Nuevo RUS "
                                     "y 4.ª categoría se calculan aparte. Lo define el asesor legal en la Etapa 2"),
    )


SIN_CONTADOR = E(0, "SUPUESTO: por validar — declaran ellos mismos en SUNAT Virtual; un contador costaría S/ 50–150/mes")

IDEAS: list[Idea] = [
    # ------------------------------------------------------------------ 1
    Idea(
        nombre="1. Chatbot WhatsApp + panel (fichas 08+09+14)",
        unidad="una mensualidad de un cliente activo",
        precio_venta=E(150, "SUPUESTO: por validar — punto medio del rango S/ 80–200/mes de las fichas 08/09 (S-10); con IGV incluido"),
        rango_precio=(80, 200),
        **tributos_uniformes(),
        inversion_inicial={
            "chip / número de WhatsApp dedicado": E(15, "ESTIMADO: ficha 08 (S/ 10–20), punto medio"),
            "dominio para el panel (1 año)": E(50, "ESTIMADO: ficha 14 (S/ 40–60)"),
            "hosting e IA durante la construcción, antes de la 1.ª venta (1,5 meses)": E(
                1.5 * (35 + 35), "ESTIMADO: 4–6 semanas de construcción (ficha 08) × los dos costos fijos mensuales de abajo"),
            "verificación de negocio en Meta": E(0, "ESTIMADO: Meta no publica tarifa por verificar el negocio; cobra por mensaje (ficha 08)"),
        },
        costos_fijos_mensuales={
            "servidor / hosting (VPS básico)": E(35, "ESTIMADO: ficha 08 (S/ 20–50/mes); ~US$ 10 × TC 3,44"),
            "plataforma de IA (cuenta, pruebas; no por cliente)": E(35, "ESTIMADO: ficha 08 (S/ 20–60/mes), punto medio bajo"),
            "contabilidad / declaraciones": SIN_CONTADOR,
        },
        costo_variable_unitario={
            "API de Meta: mensajes de plantilla por cliente/mes": E(
                100 * 0.038 * TIPO_CAMBIO.valor,
                "ESTIMADO: 100 mensajes de plantilla de utilidad por cliente/mes × US$ 0,038 (ficha 08: SleekFlow 2025) × TC 3,44; "
                "las respuestas dentro de la ventana de 24 h no pagan"),
            "IA / tokens por cliente/mes": E(15, "ESTIMADO: ~1 500 respuestas/mes × ~S/ 0,01 por respuesta con un modelo económico; sin medición real"),
            "servidor incremental por cliente/mes": E(5, "ESTIMADO: el VPS básico soporta ~10 clientes; se prorratea la ampliación"),
        },
        escenarios=escenarios_estandar(
            pes=(1, 0.0, 100 / 150),
            base=(2, 0.20, 1.0),
            opt=(3, 0.25, 1.0),
            origen={
                "pesimista": "SUPUESTO: por validar — solo el piloto con la academia de marinera (S-14) a S/ 100 (precio de contacto) "
                             "y ningún cliente nuevo en 12 meses: primera venta en la vida de los socios y canal solo por referidos",
                "base": "SUPUESTO: por validar — 2 contactos (academia, librería) pagan desde el mes 1; +20 %/mes neto de bajas, solo por referidos",
                "optimista": "SUPUESTO: por validar — 3 clientes en el mes 1 (academia, librería, galletas) y +25 %/mes neto con referidos y contenido en redes",
            }),
        horas_fijas_semana=E(12, "ESTIMADO: mantenimiento y mejoras 6 h (Carlos) + demos y venta 4 h (Indira) + administración 2 h. "
                                 "NO incluye las 80–120 h de construcción previas al mes 1 (4–6 semanas × 20 h, ficha 08)"),
        horas_por_unidad=E(2.0, "ESTIMADO: configuración inicial 4 h (ficha 08: 3–5 h) prorrateada en 4 meses de permanencia + 1 h/mes de soporte por cliente"),
        notas="Fichas 08, 09 y 14 como un producto con módulos (ventas, atención, panel de KPI). La unidad es una mensualidad; el "
              "crecimiento es neto de bajas (no se modela churn aparte). Con API no oficial el costo de Meta sería 0 pero aplica S-11 (bloqueo).",
    ),
    # ------------------------------------------------------------------ 2
    Idea(
        nombre="2. Cliente incógnito + micro-capacitación (ficha 10)",
        unidad="un servicio: visita incógnita + informe + sesión de capacitación",
        precio_venta=E(180, "SUPUESTO: por validar — dentro del rango S/ 120–250 de la ficha 10; S-09 fija el piso en S/ 100; con IGV incluido"),
        rango_precio=(120, 250),
        **tributos_uniformes(),
        inversion_inicial={
            "consumo en 3 visitas de muestra (portafolio de informes)": E(3 * 25, "ESTIMADO: ficha 10: S/ 15–30 por visita × 3 visitas a S/ 25"),
            "transporte a las 3 visitas de muestra": E(3 * 10, "ESTIMADO: S/ 10 ida y vuelta en transporte público por visita"),
            "plantilla de pauta e informe (Canva/Docs gratuitos)": E(0, "ESTIMADO: herramientas gratuitas; se redacta desde cero (S-06)"),
        },
        costos_fijos_mensuales={
            "promoción mínima en Instagram (publicaciones impulsadas)": E(40, "SUPUESTO: por validar — mínimo para llegar a dueños sin prospección en frío (F4)"),
            "herramientas digitales (Canva, Google gratuitos)": E(0, "ESTIMADO: planes gratuitos"),
            "contabilidad / declaraciones": SIN_CONTADOR,
        },
        costo_variable_unitario={
            "consumo durante la visita incógnita": E(25, "ESTIMADO: ficha 10 (S/ 15–30)"),
            "transporte: visita + sesión de capacitación": E(2 * 10, "ESTIMADO: 2 desplazamientos ida y vuelta × S/ 10 en Lima Norte"),
            "material de la sesión (impresiones)": E(5, "ESTIMADO: ~10 hojas impresas"),
        },
        escenarios=escenarios_estandar(
            pes=(1, 0.0, 0.75),
            base=(2, 0.15, 1.0),
            opt=(3, 0.20, 1.0),
            origen={
                "pesimista": "SUPUESTO: por validar — 1 servicio al mes a conocidos con 25 % de descuento y sin referidos: los dueños esperan el informe gratis (ficha 10)",
                "base": "SUPUESTO: por validar — 2 locales en el mes 1 (conocidos y locales que ya frecuentan) y +15 %/mes por referidos entre dueños",
                "optimista": "SUPUESTO: por validar — 3 en el mes 1 y +20 %/mes; cada cliente refiere a otro dueño del barrio",
            }),
        horas_fijas_semana=E(3, "ESTIMADO: venta por mensaje directo, contenido en redes y coordinación"),
        horas_por_unidad=E(6, "ESTIMADO: ficha 10: visita 1–1,5 h + informe 2 h + capacitación 1–1,5 h + coordinación 1 h"),
        notas="Nada que construir. Visitas en noches y fines de semana. El régimen real podría ser 4.ª categoría (ver variantes).",
    ),
    # ------------------------------------------------------------------ 3
    Idea(
        nombre="3. Taller 'construye tu app con IA' (ficha 13)",
        unidad="un asiento en un taller en vivo de 2–3 h",
        precio_venta=E(40, "SUPUESTO: por validar — rango S/ 30–50 de la ficha 13 (S-15); precio de Coderhouse Perú DESCONOCIDO; con IGV incluido"),
        rango_precio=(30, 50),
        **tributos_uniformes(),
        inversion_inicial={
            "micrófono USB o de solapa": E(60, "ESTIMADO: ficha 13"),
            "aro de luz / iluminación básica": E(40, "ESTIMADO: ficha 13 fija el tope en S/ 100 en total"),
        },
        costos_fijos_mensuales={
            "Zoom Pro (talleres de 2–3 h; Meet gratuito corta a 60 min)": E(15.99 * TIPO_CAMBIO.valor, "ESTIMADO: US$ 15,99/mes × TC 3,44"),
            "plan de herramienta de IA para demostraciones": E(20 * TIPO_CAMBIO.valor,
                "ESTIMADO: US$ 20/mes × TC 3,44; la ficha 13 dice que Carlos ya lo paga; se incluye igual como costo del negocio"),
            "edición de video (CapCut gratuito)": E(0, "ESTIMADO: plan gratuito"),
            "contabilidad / declaraciones": SIN_CONTADOR,
        },
        costo_variable_unitario={
            "créditos de IA usados por el asistente durante el taller": E(3, "ESTIMADO: ~3 h de uso de un modelo económico por persona; sin medición"),
            "certificado y material digital (PDF)": E(0, "ESTIMADO: sin costo"),
        },
        escenarios=escenarios_estandar(
            pes=(4, -0.15, 0.75),
            base=(10, 0.10, 1.0),
            opt=(15, 0.20, 1.0),
            origen={
                "pesimista": "SUPUESTO: por validar — 4 compañeros de la UPN con 25 % de descuento en el mes 1 y 15 % menos cada mes: sin audiencia no entra gente nueva",
                "base": "SUPUESTO: por validar — cohorte de 10 en el mes 1 (S-15: ≥ 8 interesados) y +10 %/mes por el contenido en TikTok/YouTube, sin pauta pagada",
                "optimista": "SUPUESTO: por validar — 15 asientos en el mes 1 y +20 %/mes si algún video gana tracción",
            }),
        horas_fijas_semana=E(8, "ESTIMADO: ficha 13: 3–4 videos cortos/semana 6–8 h (Carlos) + inscripciones y cobros 1–2 h (Indira)"),
        horas_por_unidad=E(7 / 12, "ESTIMADO: taller de 3 h + 4 h de preparación por cohorte de ~12 asientos = 7 h ÷ 12"),
        notas="Casi toda la carga recae en Carlos. El régimen real podría ser 4.ª categoría (ver variantes).",
    ),
    # ------------------------------------------------------------------ 4
    Idea(
        nombre="4. Sorpresas y detalles a domicilio (ficha 06)",
        unidad="un pedido (caja de detalle + envío)",
        precio_venta=E(65 + 15, "SUPUESTO: por validar — S/ 65 por el detalle (ficha 06: S/ 40–90) + S/ 15 de envío cobrado al cliente; "
                                "precios de Lima DESCONOCIDOS (Mercado Libre Perú HTTP 403); con IGV incluido"),
        rango_precio=(40 + 15, 90 + 15),
        **tributos_uniformes(),
        inversion_inicial={
            "insumos y empaque de 5 unidades de muestra (fotos)": E(5 * (26 + 4), "ESTIMADO: 5 × (insumos S/ 26 + empaque S/ 4); ficha 06: S/ 100–200"),
            "herramientas (pistola de silicona, tijeras, cúter)": E(30, "ESTIMADO: precio de bazar"),
            "fotos de catálogo": E(0, "ESTIMADO: celular propio"),
        },
        costos_fijos_mensuales={
            "promoción mínima en Instagram/TikTok (publicaciones impulsadas)": E(50, "SUPUESTO: por validar — sin audiencia previa"),
            "herramientas (Canva, WhatsApp Business)": E(0, "ESTIMADO: gratuitas"),
            "contabilidad / declaraciones": SIN_CONTADOR,
        },
        costo_variable_unitario={
            "insumos (caja, dulces/galletas de proveedor, flor o peluche pequeño, tarjeta)": E(
                5 + 10 + 8 + 3, "ESTIMADO: caja 5 + dulces 10 + flor/peluche 8 + tarjeta/globo 3; sin cotización local (S-17)"),
            "empaque (papel, cinta, relleno)": E(4, "ESTIMADO: precio de bazar"),
            "motorizado en Lima": E(15, "ESTIMADO: ficha 06 (S/ 10–20)"),
        },
        escenarios=escenarios_estandar(
            pes=(3, 0.0, 0.9),
            base=(8, 0.10, 1.0),
            opt=(15, 0.15, 1.0),
            origen={
                "pesimista": "SUPUESTO: por validar — 3 pedidos al mes de conocidos con 10 % de descuento, sin crecimiento: sin ventaja frente a cientos de cuentas de detalles",
                "base": "SUPUESTO: por validar — 8 pedidos en el mes 1 (campaña navideña, compañeros y conocidos) y +10 %/mes por redes",
                "optimista": "SUPUESTO: por validar — 15 en el mes 1 y +15 %/mes si fotos y reseñas traen clientes nuevos",
            }),
        horas_fijas_semana=E(5, "ESTIMADO: contenido en redes, atención por WhatsApp, compras de insumos"),
        horas_por_unidad=E(1.5, "ESTIMADO: armado 1 h + coordinación y entrega al motorizado 0,5 h"),
        notas="Condición F5: sin alimentos preparados en casa; comida solo de proveedores. 'SorprendeLima' pendiente de revisión por mercado.",
    ),
    # ------------------------------------------------------------------ 5
    Idea(
        nombre="5. Decoración para eventos pequeños (ficha 05)",
        unidad="un evento (arco de globos + fondo sencillo)",
        precio_venta=E(250, "SUPUESTO: por validar — arco de globos + fondo sencillo en Lima Norte; precios locales DESCONOCIDOS (ficha 05); "
                            "rango S/ 150–400; con IGV incluido"),
        rango_precio=(150, 400),
        **tributos_uniformes(),
        inversion_inicial={
            "estructura de arco / soporte de fondo (reutilizable)": E(115, "ESTIMADO: ficha 05 (S/ 80–150)"),
            "inflador eléctrico (reutilizable)": E(80, "ESTIMADO: ficha 05 (S/ 60–100)"),
            "telas y fondo (reutilizables)": E(60, "ESTIMADO: ficha 05 (globos y telas iniciales S/ 100–150, parte reutilizable)"),
            "herramientas (pistola de silicona, cinta, ganchos, nylon)": E(30, "ESTIMADO: precio de bazar"),
            "globos de 2 montajes propios para fotos (portafolio)": E(2 * 30, "ESTIMADO: 2 × S/ 30 en globos, sin transporte"),
        },
        costos_fijos_mensuales={
            "promoción mínima en Instagram/Facebook": E(40, "SUPUESTO: por validar — sin portafolio ni audiencia"),
            "almacenamiento de estructuras": E(0, "ESTIMADO: en casa (pregunta abierta de la ficha 05)"),
            "contabilidad / declaraciones": SIN_CONTADOR,
        },
        costo_variable_unitario={
            "globos (100–150 unidades) y consumibles": E(60, "ESTIMADO: sin cotización local (S-17)"),
            "extras (flores de papel, letreros, cinta)": E(20, "ESTIMADO: sin cotización local"),
            "transporte de material ida y vuelta (taxi/flete; sin auto)": E(2 * 20, "ESTIMADO: 2 × S/ 20 en Lima Norte; ficha 05: el flete entra en el costo"),
        },
        escenarios=escenarios_estandar(
            pes=(1, 0.0, 0.8),
            base=(2, 0.10, 1.0),
            opt=(3, 0.12, 1.0),
            origen={
                "pesimista": "SUPUESTO: por validar — 1 evento al mes de familiares con 20 % de descuento, sin referidos",
                "base": "SUPUESTO: por validar — 2 eventos/mes desde el mes 1 (fin de año) y +10 %/mes por fotos y referidos",
                "optimista": "SUPUESTO: por validar — 3 en el mes 1 y +12 %/mes; límite físico: solo fines de semana (F2)",
            }),
        horas_fijas_semana=E(4, "ESTIMADO: cotizaciones por WhatsApp, contenido, compras"),
        horas_por_unidad=E(10, "ESTIMADO: ficha 05: armado previo 3–4 h + montaje 2–3 h entre dos personas + desmontaje 1 h + coordinación 1 h ≈ 10 horas-persona"),
        notas="Inversión en kit reutilizable. Solo fines de semana (F2). Se adelanta el costo variable antes de cobrar salvo adelanto del 50 %.",
    ),
]


# =============================================================================
# 2b. DESEMPATE A vs B — Fase 4 de /explorar (entregables/plan-de-validacion.md, sección D)
# =============================================================================
# Métrica única: margen real por hora a 3 meses (MRH), en S/ por hora, con lo cobrado al 2026-11-15.
#   A (chatbot):  [adelantos − costos directos del piloto + 2 × (mensualidad − CV_S29) × clientes que siguen el mes 2]
#                 ÷ [horas A0+A1+A2 + horas de construcción de A3 (o 30 si no terminó) + 2 h × clientes × 2 meses]
#   B (detalles): [adelantos − costos reales de insumos, empaque y motorizado] ÷ [horas B1+B2+B3+B4]
# Regla (fijada antes de ejecutar): solo una pasó → esa; ambas → mayor MRH, y si la diferencia es < 20 % gana A
# (ingreso recurrente); ninguna → volver a Fase 3. Es la recomendación de la agencia, no un veto.
ORIGEN_DESEMPATE = "DESCONOCIDO — se llena con los registros reales de los socios al 2026-11-15"
CV_S29_CHATBOT = E(33.07, "ESTIMADO: costo variable por cliente/mes del chatbot (S-29) fijado en el plan de validación, sección D; "
                          "igual a la suma de costo_variable_unitario de la idea 1 (Meta 13,07 + IA 15 + servidor 5)")
HORAS_CONSTRUCCION_A3_TOPE = E(30, "VERIFICADO: plan de validación, sección D y criterio de A3 — si A3 no terminó se cuentan 30 h")
HORAS_POR_CLIENTE_MES_A = E(2, "VERIFICADO: plan de validación, sección D — 2 h por cliente por mes")
MESES_RECURRENCIA_A = E(2, "VERIFICADO: plan de validación, sección D — meses 2 y 3 del piloto")
UMBRAL_EMPATE = E(0.20, "VERIFICADO: plan de validación, regla 2 — diferencia < 20 % favorece A")


@dataclass
class Dato:
    """Un dato real del desempate. `valor` None = todavía no registrado; el comando lo lista como faltante."""

    valor: float | bool | None
    origen: str
    descripcion: str


def D(descripcion: str, valor=None, origen: str = ORIGEN_DESEMPATE) -> Dato:
    return Dato(valor, origen, descripcion)


DATOS_DESEMPATE: dict[str, dict[str, Dato]] = {
    "A": {
        "paso_A2": D("¿A pasó su criterio A2? (≥ 1 pago de S/ 100 con inicio ≤ 2026-11-30) — sí/no"),
        "adelantos_cobrados": D("S/ adelantos cobrados por pilotos (Yape/transferencia con comprobante)"),
        "costos_directos_piloto": D("S/ costos directos reales del piloto: chip, Meta, IA, hosting (recibos)"),
        "mensualidad_acordada": D("S/ mensualidad acordada por escrito para los meses 2 y 3"),
        "clientes_siguen_mes2": D("n.º de clientes que confirman por WhatsApp seguir el mes 2"),
        "horas_A0_A1_A2": D("h registradas de A0 + A1 + A2, ambos socios"),
        "a3_terminado": D("¿A3 (piloto operativo) terminó al cierre del registro? — sí/no"),
        "horas_construccion_A3": D("h cronometradas de construcción en A3 (solo si a3_terminado = sí; si no, se usan 30 h)"),
    },
    "B": {
        "paso_B2": D("¿B pasó su criterio B2? (≥ 6 pedidos pagados, ≥ S/ 480, ≥ 2 fuera del círculo, ≥ 3/5 sin descuento) — sí/no"),
        "adelantos_cobrados": D("S/ adelantos cobrados por pedidos (incluye un 2.º pedido pagado para diciembre si lo hubo)"),
        "costos_reales_pedidos": D("S/ costos reales de insumos, empaque y motorizado de esos pedidos (boletas)"),
        "horas_B1_B2_B3_B4": D("h registradas de B1 + B2 + B3 + B4, ambos socios"),
    },
}

# Cifras ILUSTRATIVAS del plan (sección D y tabla de horas) — solo para fijar escala y probar la función. NO son datos.
ORIGEN_ILUSTRATIVO = "ILUSTRATIVO: plan de validación 2026-10-22, sección D y tabla resumen; no es un dato real"
DATOS_DESEMPATE_EJEMPLO: dict[str, dict[str, Dato]] = {
    "A": {
        "paso_A2": D("criterio A2", True, ORIGEN_ILUSTRATIVO),
        "adelantos_cobrados": D("1 piloto × S/ 100", 100, ORIGEN_ILUSTRATIVO),
        "costos_directos_piloto": D("chip S/ 15 (tabla Dinero) + Meta/IA/hosting del mes del piloto S/ 33,07 (S-29)",
                                    15 + 33.07, ORIGEN_ILUSTRATIVO),
        "mensualidad_acordada": D("S/ 100 (precio de contacto, escenario pesimista de la idea 1)", 100, ORIGEN_ILUSTRATIVO),
        "clientes_siguen_mes2": D("1 cliente", 1, ORIGEN_ILUSTRATIVO),
        "horas_A0_A1_A2": D("tabla resumen: A0 2+0, A1 5+6, A2 9+4,5", 2 + 0 + 5 + 6 + 9 + 4.5, ORIGEN_ILUSTRATIVO),
        "a3_terminado": D("se asume que no terminó → 30 h", False, ORIGEN_ILUSTRATIVO),
        "horas_construccion_A3": D("no aplica (30 h por regla)", None, ORIGEN_ILUSTRATIVO),
    },
    "B": {
        "paso_B2": D("criterio B2", True, ORIGEN_ILUSTRATIVO),
        "adelantos_cobrados": D("6 pedidos × S/ 80", 6 * 80, ORIGEN_ILUSTRATIVO),
        "costos_reales_pedidos": D("6 pedidos × S/ 40", 6 * 40, ORIGEN_ILUSTRATIVO),
        "horas_B1_B2_B3_B4": D("tabla resumen: B1 2+4, B2 1+6,5, B3 0,5+2, B4 2+7", 2 + 4 + 1 + 6.5 + 0.5 + 2 + 2 + 7, ORIGEN_ILUSTRATIVO),
    },
}
# Lo que el plan dice que deberían dar las cifras ilustrativas (sección D): A 1 piloto ≈ 6, A 2 pilotos ≈ 10, B ≈ 10 S/ por hora.
MRH_PLAN_ILUSTRATIVO = {"A_1_piloto": 6.0, "A_2_pilotos": 10.0, "B": 10.0}


# =============================================================================
# 3. VALIDACIÓN
# =============================================================================


def validar(ideas: list[Idea]) -> list[str]:
    errores = []
    globales = {"CAPITAL_INICIAL": CAPITAL_INICIAL, "APORTE_MENSUAL": APORTE_MENSUAL, "TOPE_INVERSION_F1": TOPE_INVERSION_F1,
                "TOPE_FIJOS_F1": TOPE_FIJOS_F1, "TOPE_HORAS_F2": TOPE_HORAS_F2, "TIPO_CAMBIO": TIPO_CAMBIO, "RMV": RMV,
                "SUELDO_OBJETIVO": SUELDO_OBJETIVO, "CUOTA_RUS_CAT1": CUOTA_RUS_CAT1, "TASA_RENTA_4TA": TASA_RENTA_4TA,
                "PASARELA_ALT_TASA": PASARELA_ALT_TASA, "PASARELA_ALT_FIJO": PASARELA_ALT_FIJO}
    for ruta, e in globales.items():
        if e.etiqueta() == "EJEMPLO" and not ES_EJEMPLO:
            errores.append(f"[globales] {ruta}: origen EJEMPLO con ES_EJEMPLO = False")
        elif e.etiqueta() not in ETIQUETAS_VALIDAS + ("EJEMPLO",):
            errores.append(f"[globales] {ruta}: origen inválido '{e.origen}'")
    for idea in ideas:
        for ruta, e in idea.entradas():
            etq = e.etiqueta()
            if etq == "EJEMPLO":
                if not ES_EJEMPLO:
                    errores.append(f"[{idea.nombre}] {ruta}: origen EJEMPLO con ES_EJEMPLO = False")
            elif etq not in ETIQUETAS_VALIDAS:
                errores.append(f"[{idea.nombre}] {ruta}: origen inválido '{e.origen}' "
                               f"(usar VERIFICADO/ESTIMADO/SUPUESTO/DESCONOCIDO)")
        if set(idea.escenarios) != {"pesimista", "base", "optimista"}:
            errores.append(f"[{idea.nombre}] los escenarios deben ser pesimista, base y optimista")
    return errores


# =============================================================================
# 4. CÁLCULOS
# =============================================================================


def economia_unitaria(idea: Idea, factor_precio: float = 1.0, *, igv_aplica: bool | None = None,
                      tasa_renta: float | None = None, tasa_pasarela: float | None = None,
                      fijo_pasarela: float | None = None, fijos_extra: float = 0.0) -> dict:
    """Economía por unidad. Los argumentos con nombre permiten variantes (régimen, pasarela)."""
    igv_aplica = idea.igv_aplica if igv_aplica is None else igv_aplica
    tasa_renta = idea.tasa_impuesto_renta.valor if tasa_renta is None else tasa_renta
    tasa_pasarela = idea.tasa_pasarela.valor if tasa_pasarela is None else tasa_pasarela
    fijo_pasarela = idea.fijo_pasarela.valor if fijo_pasarela is None else fijo_pasarela
    precio = idea.precio_venta.valor * factor_precio
    ingreso_neto = precio / (1 + idea.tasa_igv.valor) if igv_aplica else precio
    comision = precio * tasa_pasarela + fijo_pasarela
    impuesto = ingreso_neto * tasa_renta
    cv_directo = sum(e.valor for e in idea.costo_variable_unitario.values())
    cv_total = cv_directo + comision + impuesto
    margen = ingreso_neto - cv_total
    fijos = sum(e.valor for e in idea.costos_fijos_mensuales.values()) + fijos_extra
    pe_unid = fijos / margen if margen > 0 else math.inf
    return {
        "precio": precio,
        "ingreso_neto_unitario": ingreso_neto,
        "comision_pasarela": comision,
        "impuesto_renta_unitario": impuesto,
        "costo_variable_directo": cv_directo,
        "costo_variable_total": cv_total,
        "margen_unitario": margen,
        "margen_pct_sobre_ingreso_neto": (margen / ingreso_neto) if ingreso_neto else 0.0,
        "costos_fijos_mensuales": fijos,
        "punto_equilibrio_unidades_mes": pe_unid,
        "punto_equilibrio_soles_mes": pe_unid * precio if math.isfinite(pe_unid) else math.inf,
    }


def inversion_total(idea: Idea) -> float:
    return sum(e.valor for e in idea.inversion_inicial.values())


def flujo_12_meses(idea: Idea, esc: Escenario) -> dict:
    eu = economia_unitaria(idea, esc.factor_precio.valor)
    inv = inversion_total(idea)
    meses, acum, recuperacion = [], -inv, None
    mes_sueldo_1, mes_sueldo_2 = None, None
    financiable = inv <= CAPITAL_INICIAL.valor
    for m in range(1, MESES + 1):
        u = unidades_en_mes(esc, m)
        ingresos = u * eu["ingreso_neto_unitario"]
        variables = u * eu["costo_variable_total"]
        neto = ingresos - variables - eu["costos_fijos_mensuales"]
        acum += neto
        if recuperacion is None and acum >= 0:
            recuperacion = m
        if mes_sueldo_1 is None and neto >= SUELDO_OBJETIVO.valor:
            mes_sueldo_1 = m
        if mes_sueldo_2 is None and neto >= 2 * SUELDO_OBJETIVO.valor:
            mes_sueldo_2 = m
        # el capital disponible hasta el mes m es el inicial más los aportes de m meses
        if -acum > CAPITAL_INICIAL.valor + APORTE_MENSUAL.valor * m:
            financiable = False
        meses.append({"mes": m, "unidades": u, "ingresos_netos": ingresos,
                      "costos_variables": variables, "costos_fijos": eu["costos_fijos_mensuales"],
                      "flujo_neto": neto, "acumulado": acum,
                      "horas_semana": idea.horas_semana_en_mes(esc, m)})
    return {
        "escenario": esc.nombre,
        "economia_unitaria": eu,
        "inversion_inicial": inv,
        "meses": meses,
        "unidades_12m": sum(x["unidades"] for x in meses),
        "unidades_mes12": meses[-1]["unidades"],
        "ingresos_netos_12m": sum(x["ingresos_netos"] for x in meses),
        "resultado_12m_tras_inversion": acum,
        "flujo_neto_mes12": meses[-1]["flujo_neto"],
        "mes_recuperacion": recuperacion,  # None = no se recupera en 12 meses
        "peor_acumulado": min(x["acumulado"] for x in meses + [{"acumulado": -inv}]),
        "financiable_con_capital_real": financiable,
        "mes_cubre_1_sueldo": mes_sueldo_1,  # flujo neto mensual ≥ SUELDO_OBJETIVO
        "mes_cubre_2_sueldos": mes_sueldo_2,
        "horas_semana_mes12": meses[-1]["horas_semana"],
        "mes_supera_F2": next((x["mes"] for x in meses if x["horas_semana"] > TOPE_HORAS_F2.valor), None),
    }


def resultado_12m(idea: Idea, esc: Escenario, **override) -> float:
    """Resultado acumulado a 12 meses (tras inversión) con entradas sobreescritas.
    Claves de override: precio, cv_directo, fijos, inversion, unidades_mes1, crecimiento,
    tasa_pasarela, tasa_impuesto_renta. Se usa para la sensibilidad y el rango de precio."""
    precio = override.get("precio", idea.precio_venta.valor) * esc.factor_precio.valor
    ingreso_neto = precio / (1 + idea.tasa_igv.valor) if idea.igv_aplica else precio
    comision = precio * override.get("tasa_pasarela", idea.tasa_pasarela.valor) + idea.fijo_pasarela.valor
    impuesto = ingreso_neto * override.get("tasa_impuesto_renta", idea.tasa_impuesto_renta.valor)
    cv_directo = override.get("cv_directo", sum(e.valor for e in idea.costo_variable_unitario.values()))
    margen = ingreso_neto - (cv_directo + comision + impuesto)
    fijos = override.get("fijos", sum(e.valor for e in idea.costos_fijos_mensuales.values()))
    inv = override.get("inversion", inversion_total(idea))
    u1 = override.get("unidades_mes1", esc.unidades_mes1.valor)
    g = override.get("crecimiento", esc.crecimiento_mensual.valor)
    unidades = MESES * u1 if g == 0 else u1 * ((1 + g) ** MESES - 1) / g
    return unidades * margen - MESES * fijos - inv


def sensibilidad(idea: Idea, variacion: float = 0.20) -> list[dict]:
    """Una variable a la vez, ±variacion, sobre el escenario base. Ordenada por impacto."""
    esc = idea.escenarios["base"]
    base = resultado_12m(idea, esc)
    variables = {
        "precio": idea.precio_venta.valor,
        "unidades_mes1": esc.unidades_mes1.valor,
        "crecimiento": esc.crecimiento_mensual.valor,
        "cv_directo": sum(e.valor for e in idea.costo_variable_unitario.values()),
        "fijos": sum(e.valor for e in idea.costos_fijos_mensuales.values()),
        "inversion": inversion_total(idea),
        "tasa_pasarela": idea.tasa_pasarela.valor,
        "tasa_impuesto_renta": idea.tasa_impuesto_renta.valor,
    }
    filas = []
    for k, v in variables.items():
        menos = resultado_12m(idea, esc, **{k: v * (1 - variacion)})
        mas = resultado_12m(idea, esc, **{k: v * (1 + variacion)})
        filas.append({"variable": k, "valor_base": v, "resultado_base": base,
                      f"resultado_-{int(variacion*100)}%": menos,
                      f"resultado_+{int(variacion*100)}%": mas,
                      "rango": abs(mas - menos)})
    filas.sort(key=lambda f: f["rango"], reverse=True)
    return filas


def rango_precio(idea: Idea) -> list[dict] | None:
    """Resultado en los extremos del rango de precio (cuando el precio es un supuesto con rango)."""
    if not idea.rango_precio:
        return None
    esc = idea.escenarios["base"]
    out = []
    for p in idea.rango_precio:
        eu = economia_unitaria(idea, p / idea.precio_venta.valor)
        out.append({"precio": p, "margen_unitario": eu["margen_unitario"],
                    "margen_pct": eu["margen_pct_sobre_ingreso_neto"],
                    "pe_unidades_mes": eu["punto_equilibrio_unidades_mes"],
                    "resultado_12m_base": resultado_12m(idea, esc, precio=p)})
    return out


def variantes(idea: Idea) -> dict:
    """Mismo precio de lista bajo otros regímenes o con pasarela con comisión. Solo economía unitaria."""
    return {
        "RER (supuesto base: IGV 18 % + renta 1,5 %)": economia_unitaria(idea),
        "Nuevo RUS cat. 1 (sin IGV, sin renta %, cuota S/ 20/mes como fijo; solo boletas)":
            economia_unitaria(idea, igv_aplica=False, tasa_renta=0.0, fijos_extra=CUOTA_RUS_CAT1.valor),
        "4.ª categoría (sin IGV, retención 8 %; solo servicios personales: ideas 2 y 3)":
            economia_unitaria(idea, igv_aplica=False, tasa_renta=TASA_RENTA_4TA.valor),
        "RER + pasarela Culqi (3,44 % + US$ 0,20)":
            economia_unitaria(idea, tasa_pasarela=PASARELA_ALT_TASA.valor, fijo_pasarela=PASARELA_ALT_FIJO.valor),
    }


def unidades_para_sueldo(idea: Idea, n_sueldos: int) -> float:
    eu = economia_unitaria(idea)
    if eu["margen_unitario"] <= 0:
        return math.inf
    return (eu["costos_fijos_mensuales"] + n_sueldos * SUELDO_OBJETIVO.valor) / eu["margen_unitario"]


def evaluar(idea: Idea) -> dict:
    inv = inversion_total(idea)
    eu = economia_unitaria(idea)
    return {
        "nombre": idea.nombre,
        "unidad": idea.unidad,
        "inversion_inicial": inv,
        "economia_unitaria": eu,
        "escenarios": {n: flujo_12_meses(idea, esc) for n, esc in idea.escenarios.items()},
        "sensibilidad": sensibilidad(idea),
        "rango_precio": rango_precio(idea),
        "variantes": variantes(idea),
        "horas_semana_requeridas": idea.horas_semana_requeridas.valor,
        "horas_fijas_semana": idea.horas_fijas_semana.valor,
        "horas_por_unidad": idea.horas_por_unidad.valor,
        "f1_inversion_ok": inv <= TOPE_INVERSION_F1.valor,
        "f1_fijos_ok": eu["costos_fijos_mensuales"] <= TOPE_FIJOS_F1.valor,
        "f2_horas_ok_base": idea.horas_semana_requeridas.valor <= TOPE_HORAS_F2.valor,
        "unidades_mes_para_1_sueldo": unidades_para_sueldo(idea, 1),
        "unidades_mes_para_2_sueldos": unidades_para_sueldo(idea, 2),
        "horas_semana_para_1_sueldo": idea.horas_fijas_semana.valor + idea.horas_por_unidad.valor * unidades_para_sueldo(idea, 1) / SEMANAS_POR_MES,
        "horas_semana_para_2_sueldos": idea.horas_fijas_semana.valor + idea.horas_por_unidad.valor * unidades_para_sueldo(idea, 2) / SEMANAS_POR_MES,
    }


# =============================================================================
# 5. SALIDA EN TEXTO
# =============================================================================


def s(x: float) -> str:
    if x is None:
        return "—"
    if isinstance(x, float) and not math.isfinite(x):
        return "no alcanzable"
    return f"S/ {x:,.0f}" if abs(x) >= 100 else f"S/ {x:,.2f}"


def n(x: float, d: int = 1) -> str:
    if x is None:
        return "—"
    if isinstance(x, float) and not math.isfinite(x):
        return "n/a"
    return f"{x:,.{d}f}"


def banner() -> str:
    if ES_EJEMPLO:
        return ("\n" + "!" * 78 + "\n!!  VALORES DE EJEMPLO. No describen ningún negocio real.\n"
                "!!  Borrar las ideas de ejemplo y poner ES_EJEMPLO = False antes del análisis real.\n"
                + "!" * 78 + "\n")
    return ""


def imprimir_globales() -> None:
    print("Parámetros globales (mismos para todas las ideas):")
    for nombre, e in (("capital inicial", CAPITAL_INICIAL), ("aporte mensual máximo", APORTE_MENSUAL),
                      ("tope inversión F1", TOPE_INVERSION_F1), ("tope fijos F1", TOPE_FIJOS_F1),
                      ("tope horas F2", TOPE_HORAS_F2), ("tipo de cambio S/ por US$", TIPO_CAMBIO), ("RMV", RMV),
                      ("sueldo objetivo por socio", SUELDO_OBJETIVO), ("cuota Nuevo RUS cat. 1", CUOTA_RUS_CAT1),
                      ("retención 4.ª categoría", TASA_RENTA_4TA), ("pasarela alternativa %", PASARELA_ALT_TASA),
                      ("pasarela alternativa fijo S/", PASARELA_ALT_FIJO)):
        print(f"  {nombre:<32} {e.valor:>12,.4g}  [{e.origen}]")


def imprimir_idea(idea: Idea, r: dict) -> None:
    eu = r["economia_unitaria"]
    print(f"\n=== {idea.nombre} ===")
    print(f"Unidad: {idea.unidad}.")
    print(f"Horas/semana requeridas (pico mes 12, base): {r['horas_semana_requeridas']:.1f} "
          f"[{idea.horas_semana_requeridas.origen}]")
    if idea.notas:
        print(f"Notas: {idea.notas}")
    print("\nEntradas (valor · origen):")
    for ruta, e in idea.entradas():
        print(f"  {ruta:<70} {e.valor:>12,.4g}  [{e.origen}]")
    print("\nEconomía unitaria (precio de lista):")
    print(f"  Precio al cliente              {s(eu['precio'])}")
    print(f"  Ingreso neto (sin IGV)         {s(eu['ingreso_neto_unitario'])}")
    print(f"  Comisión pasarela              {s(eu['comision_pasarela'])}")
    print(f"  Impuesto a la renta            {s(eu['impuesto_renta_unitario'])}")
    print(f"  Costo variable directo         {s(eu['costo_variable_directo'])}")
    print(f"  Costo variable total           {s(eu['costo_variable_total'])}")
    print(f"  Margen por unidad              {s(eu['margen_unitario'])}  "
          f"({eu['margen_pct_sobre_ingreso_neto']*100:.1f} % del ingreso neto)")
    print(f"  Costos fijos mensuales         {s(eu['costos_fijos_mensuales'])}")
    pe = eu["punto_equilibrio_unidades_mes"]
    print(f"  Punto de equilibrio            {pe:.1f} unidades/mes = {s(eu['punto_equilibrio_soles_mes'])} de ventas/mes" if math.isfinite(pe)
          else "  Punto de equilibrio            no alcanzable (margen <= 0)")
    print(f"  Inversión inicial              {s(r['inversion_inicial'])}")
    print(f"  Filtro F1: inversión ≤ {s(TOPE_INVERSION_F1.valor)}: {'sí' if r['f1_inversion_ok'] else 'NO'}; "
          f"fijos ≤ {s(TOPE_FIJOS_F1.valor)}/mes: {'sí' if r['f1_fijos_ok'] else 'NO'}")
    print(f"  Para cubrir 1 sueldo ({s(SUELDO_OBJETIVO.valor)}/mes): {n(r['unidades_mes_para_1_sueldo'])} unidades/mes "
          f"≈ {n(r['horas_semana_para_1_sueldo'])} h/semana; 2 sueldos: {n(r['unidades_mes_para_2_sueldos'])} unidades/mes "
          f"≈ {n(r['horas_semana_para_2_sueldos'])} h/semana (tope F2: {TOPE_HORAS_F2.valor:.0f} h)")
    print("\nFlujo de caja a 12 meses (escenarios, no predicciones):")
    print(f"  {'escenario':<10} {'unid. 12m':>10} {'unid. m12':>10} {'ingresos 12m':>13} {'resultado 12m':>14} "
          f"{'recupera':>10} {'peor caja':>11} {'cabe cap.':>9} {'1 sueldo':>9} {'h/sem m12':>10} {'>F2 en':>7}")
    for nom, f in r["escenarios"].items():
        rec = f"mes {f['mes_recuperacion']}" if f["mes_recuperacion"] else "no en 12m"
        su = f"mes {f['mes_cubre_1_sueldo']}" if f["mes_cubre_1_sueldo"] else "no"
        f2 = f"mes {f['mes_supera_F2']}" if f["mes_supera_F2"] else "nunca"
        print(f"  {nom:<10} {f['unidades_12m']:>10,.0f} {f['unidades_mes12']:>10,.1f} {s(f['ingresos_netos_12m']):>13} "
              f"{s(f['resultado_12m_tras_inversion']):>14} {rec:>10} {s(f['peor_acumulado']):>11} "
              f"{('sí' if f['financiable_con_capital_real'] else 'NO'):>9} {su:>9} {f['horas_semana_mes12']:>10.1f} {f2:>7}")
    print("\nFlujo mensual — escenario base:")
    print(f"  {'mes':>3} {'unid.':>8} {'ingresos':>12} {'variables':>12} {'fijos':>10} {'neto':>12} {'acum.':>12} {'h/sem':>7}")
    for m in r["escenarios"]["base"]["meses"]:
        print(f"  {m['mes']:>3} {m['unidades']:>8,.1f} {s(m['ingresos_netos']):>12} "
              f"{s(m['costos_variables']):>12} {s(m['costos_fijos']):>10} {s(m['flujo_neto']):>12} "
              f"{s(m['acumulado']):>12} {m['horas_semana']:>7.1f}")
    print("\nSensibilidad (escenario base, ±20 % una variable a la vez, ordenada por impacto):")
    for f in r["sensibilidad"]:
        print(f"  {f['variable']:<20} base {s(f['resultado_base']):>12}   "
              f"-20 % → {s(f['resultado_-20%']):>12}   +20 % → {s(f['resultado_+20%']):>12}   "
              f"rango {s(f['rango'])}")
    print(f"  → La entrada que más mueve el resultado es: {r['sensibilidad'][0]['variable']}")
    if r["rango_precio"]:
        print("\nRango del precio (supuesto por validar) — extremos, escenario base:")
        for x in r["rango_precio"]:
            print(f"  precio {s(x['precio']):>9}: margen {s(x['margen_unitario']):>9} ({x['margen_pct']*100:.1f} %), "
                  f"PE {n(x['pe_unidades_mes'])} u/mes, resultado 12m base {s(x['resultado_12m_base'])}")
    print("\nVariantes (mismo precio de lista; cambia solo régimen o pasarela):")
    for nom, v in r["variantes"].items():
        print(f"  {nom:<88} margen {s(v['margen_unitario']):>9} ({v['margen_pct_sobre_ingreso_neto']*100:.1f} %), "
              f"PE {n(v['punto_equilibrio_unidades_mes'])} u/mes")


def imprimir_comparacion(ideas: list[Idea], resultados: list[dict]) -> None:
    print("\nComparación con el mismo método (escenario base salvo donde se indica):")
    print(f"  {'idea':<48} {'inversión':>9} {'margen/u':>9} {'marg%':>6} {'PE u/mes':>8} {'recup.':>9} "
          f"{'res.12m pes':>12} {'res.12m base':>13} {'res.12m opt':>12} {'h/sem':>6} {'F1':>5}")
    for idea, r in zip(ideas, resultados):
        eu, e = r["economia_unitaria"], r["escenarios"]
        pe = eu["punto_equilibrio_unidades_mes"]
        rec = f"mes {e['base']['mes_recuperacion']}" if e["base"]["mes_recuperacion"] else "no 12m"
        f1 = "sí" if (r["f1_inversion_ok"] and r["f1_fijos_ok"]) else "NO"
        print(f"  {idea.nombre[:48]:<48} {s(r['inversion_inicial']):>9} {s(eu['margen_unitario']):>9} "
              f"{eu['margen_pct_sobre_ingreso_neto']*100:>5.1f}% {n(pe):>8} {rec:>9} "
              f"{s(e['pesimista']['resultado_12m_tras_inversion']):>12} {s(e['base']['resultado_12m_tras_inversion']):>13} "
              f"{s(e['optimista']['resultado_12m_tras_inversion']):>12} {r['horas_semana_requeridas']:>6.1f} {f1:>5}")
    print("\n  Sueldo (flujo neto mensual ≥ S/ {:,.0f} por socio) y horas:".format(SUELDO_OBJETIVO.valor))
    print(f"  {'idea':<48} {'u/mes 1 sueldo':>14} {'h/sem':>6} {'u/mes 2 sueldos':>15} {'h/sem':>6} "
          f"{'mes 1 sueldo pes/base/opt':>26} {'mes 2 sueldos base/opt':>23}")
    for idea, r in zip(ideas, resultados):
        e = r["escenarios"]
        m1 = "/".join(str(e[k]["mes_cubre_1_sueldo"]) if e[k]["mes_cubre_1_sueldo"] else "no" for k in ("pesimista", "base", "optimista"))
        m2 = "/".join(str(e[k]["mes_cubre_2_sueldos"]) if e[k]["mes_cubre_2_sueldos"] else "no" for k in ("base", "optimista"))
        print(f"  {idea.nombre[:48]:<48} {n(r['unidades_mes_para_1_sueldo']):>14} {n(r['horas_semana_para_1_sueldo']):>6} "
              f"{n(r['unidades_mes_para_2_sueldos']):>15} {n(r['horas_semana_para_2_sueldos']):>6} {m1:>26} {m2:>23}")
    print("\n  Caja y capital (peor acumulado por escenario; cabe en S/ {:,.0f} + S/ {:,.0f}/mes):".format(
        CAPITAL_INICIAL.valor, APORTE_MENSUAL.valor))
    print(f"  {'idea':<48} {'peor pes':>10} {'cabe':>5} {'peor base':>10} {'cabe':>5} {'peor opt':>10} {'cabe':>5} {'>F2 base':>9} {'>F2 opt':>8}")
    for idea, r in zip(ideas, resultados):
        e = r["escenarios"]
        cols = []
        for k in ("pesimista", "base", "optimista"):
            cols += [s(e[k]["peor_acumulado"]), "sí" if e[k]["financiable_con_capital_real"] else "NO"]
        f2b = f"mes {e['base']['mes_supera_F2']}" if e["base"]["mes_supera_F2"] else "nunca"
        f2o = f"mes {e['optimista']['mes_supera_F2']}" if e["optimista"]["mes_supera_F2"] else "nunca"
        print(f"  {idea.nombre[:48]:<48} {cols[0]:>10} {cols[1]:>5} {cols[2]:>10} {cols[3]:>5} {cols[4]:>10} {cols[5]:>5} {f2b:>9} {f2o:>8}")
    print("\n  Supuesto más frágil por idea (mayor sensibilidad ±20 %, escenario base) y segundo:")
    for idea, r in zip(ideas, resultados):
        sb = r["sensibilidad"]
        print(f"    {idea.nombre[:48]:<48} 1.º {sb[0]['variable']} (rango {s(sb[0]['rango'])}) · 2.º {sb[1]['variable']} (rango {s(sb[1]['rango'])})")


# =============================================================================
# 6. HOJA DE CÁLCULO CON FÓRMULAS
# =============================================================================


def exportar_xlsx(ideas: list[Idea], ruta: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    negrita = Font(bold=True)
    gris = PatternFill("solid", fgColor="EEEEEE")
    amarillo = PatternFill("solid", fgColor="FFF2CC")

    # --- LEEME -------------------------------------------------------------
    ws = wb.active
    ws.title = "LEEME"
    filas = [
        ["Modelo financiero — agencia de creación de negocios — Fase 3 (2026-10-21)"],
        ["Generado por finanzas/modelo.py --xlsx. Las celdas amarillas son entradas; el resto son fórmulas."],
        ["Para cambiar un supuesto, edite la celda amarilla en la hoja 'Supuestos' de la idea o en 'Globales'; todo se recalcula."],
        ["Etiquetas de origen: VERIFICADO: fuente, fecha · ESTIMADO: método · SUPUESTO: por validar · DESCONOCIDO."],
        ["Si cambia un valor aquí, replíquelo en modelo.py (o al revés) para que ambos coincidan."],
        ["Las proyecciones son escenarios, no predicciones. Régimen tributario uniforme SUPUESTO (RER) hasta el informe legal."],
        [""],
    ]
    if ES_EJEMPLO:
        filas.insert(0, ["¡¡ VALORES DE EJEMPLO !! No describen ningún negocio real. Borrar antes del análisis real."])
    for f in filas:
        ws.append(f)
    ws["A1"].font = Font(bold=True, size=13, color="C00000" if ES_EJEMPLO else "000000")
    ws.column_dimensions["A"].width = 110

    # --- Globales ------------------------------------------------------------
    wsG = wb.create_sheet("Globales")
    wsG.append(["Parámetro", "Valor", "Origen"]); wsG["A1"].font = negrita
    G: dict[str, str] = {}
    for clave, nombre, e in (("capital", "Capital inicial (S/)", CAPITAL_INICIAL), ("aporte", "Aporte mensual máximo (S/)", APORTE_MENSUAL),
                             ("tope_inv", "Tope inversión F1 (S/)", TOPE_INVERSION_F1), ("tope_fijos", "Tope fijos F1 (S/ por mes)", TOPE_FIJOS_F1),
                             ("tope_horas", "Tope horas F2 (h/semana)", TOPE_HORAS_F2), ("tc", "Tipo de cambio (S/ por US$)", TIPO_CAMBIO),
                             ("rmv", "RMV (S/)", RMV), ("sueldo", "Sueldo objetivo por socio (S/ por mes)", SUELDO_OBJETIVO),
                             ("rus", "Cuota Nuevo RUS cat. 1 (S/ por mes)", CUOTA_RUS_CAT1), ("t4", "Retención 4.ª categoría", TASA_RENTA_4TA),
                             ("pas_t", "Pasarela alternativa %", PASARELA_ALT_TASA), ("pas_f", "Pasarela alternativa fijo (S/)", PASARELA_ALT_FIJO)):
        wsG.append([nombre, e.valor, e.origen]); wsG[f"B{wsG.max_row}"].fill = amarillo
        G[clave] = f"Globales!$B${wsG.max_row}"
    wsG.column_dimensions["A"].width = 40; wsG.column_dimensions["B"].width = 14; wsG.column_dimensions["C"].width = 120

    comp_refs = []
    for idx, idea in enumerate(ideas, start=1):
        p = f"I{idx}"  # prefijo corto de hoja
        # --- Supuestos -----------------------------------------------------
        wsS = wb.create_sheet(f"{p} Supuestos")
        wsS.append([f"Idea: {idea.nombre}"]); wsS["A1"].font = negrita
        wsS.append([f"Unidad: {idea.unidad}"])
        wsS.append(["Concepto", "Valor", "Origen", "Nota"]); wsS["A3"].font = negrita
        ref: dict[str, str] = {}

        def fila(clave, concepto, valor, origen, nota="", editable=True):
            wsS.append([concepto, valor, origen, nota])
            r_ = wsS.max_row
            if editable:
                wsS[f"B{r_}"].fill = amarillo
            ref[clave] = f"'{p} Supuestos'!$B${r_}"

        fila("precio", "Precio de venta al cliente (S/)", idea.precio_venta.valor, idea.precio_venta.origen, idea.precio_venta.nota)
        fila("igv_aplica", "¿Precio incluye IGV? (1 = sí, 0 = no)", 1 if idea.igv_aplica else 0, "regla", "")
        fila("tasa_igv", "Tasa IGV", idea.tasa_igv.valor, idea.tasa_igv.origen, idea.tasa_igv.nota)
        fila("tasa_pasarela", "Comisión pasarela (%)", idea.tasa_pasarela.valor, idea.tasa_pasarela.origen, idea.tasa_pasarela.nota)
        fila("fijo_pasarela", "Comisión pasarela fija por transacción (S/)", idea.fijo_pasarela.valor, idea.fijo_pasarela.origen, idea.fijo_pasarela.nota)
        fila("tasa_renta", "Impuesto a la renta (% sobre ingreso neto)", idea.tasa_impuesto_renta.valor, idea.tasa_impuesto_renta.origen, idea.tasa_impuesto_renta.nota)
        fila("h_fijas", "Horas fijas por semana", idea.horas_fijas_semana.valor, idea.horas_fijas_semana.origen, "")
        fila("h_unidad", "Horas por unidad vendida", idea.horas_por_unidad.valor, idea.horas_por_unidad.origen, "")

        def grupo(clave, titulo, items):
            wsS.append([]); wsS.append([titulo]); wsS[f"A{wsS.max_row}"].font = negrita
            inicio = wsS.max_row + 1
            for k, e in items.items():
                wsS.append([k, e.valor, e.origen, e.nota]); wsS[f"B{wsS.max_row}"].fill = amarillo
            fin = wsS.max_row
            wsS.append([f"TOTAL {titulo.lower()}", f"=SUM(B{inicio}:B{fin})", "fórmula", ""])
            wsS[f"A{wsS.max_row}"].font = negrita
            ref[clave] = f"'{p} Supuestos'!$B${wsS.max_row}"

        grupo("inversion", "Inversión inicial (S/)", idea.inversion_inicial)
        grupo("fijos", "Costos fijos mensuales (S/)", idea.costos_fijos_mensuales)
        grupo("cv_directo", "Costo variable directo por unidad (S/)", idea.costo_variable_unitario)

        wsS.append([]); wsS.append(["Escenarios"]); wsS[f"A{wsS.max_row}"].font = negrita
        wsS.append(["Escenario", "Unidades mes 1", "Crecimiento mensual", "Factor precio", "Origen"])
        for nm, esc in idea.escenarios.items():
            wsS.append([nm, esc.unidades_mes1.valor, esc.crecimiento_mensual.valor, esc.factor_precio.valor,
                        esc.unidades_mes1.origen])
            r_ = wsS.max_row
            for c in "BCD":
                wsS[f"{c}{r_}"].fill = amarillo
            ref[f"{nm}.u1"] = f"'{p} Supuestos'!$B${r_}"
            ref[f"{nm}.g"] = f"'{p} Supuestos'!$C${r_}"
            ref[f"{nm}.f"] = f"'{p} Supuestos'!$D${r_}"
        wsS.append([])
        fila("horas", "Horas/semana requeridas (pico: mes 12 del escenario base)",
             f"={ref['h_fijas']}+{ref['h_unidad']}*{ref['base.u1']}*(1+{ref['base.g']})^({MESES}-1)/({SEMANAS_POR_MES})",
             idea.horas_semana_requeridas.origen, "", editable=False)
        if idea.rango_precio:
            wsS.append(["Rango del precio por validar (mín – máx)", idea.rango_precio[0], idea.rango_precio[1], "ver hoja Unitario"])
            r_ = wsS.max_row
            ref["precio_min"] = f"'{p} Supuestos'!$B${r_}"; ref["precio_max"] = f"'{p} Supuestos'!$C${r_}"
            wsS[f"B{r_}"].fill = amarillo; wsS[f"C{r_}"].fill = amarillo
        for col, w in (("A", 60), ("B", 14), ("C", 36), ("D", 40), ("E", 14)):
            wsS.column_dimensions[col].width = w

        # --- Unitario (precio de lista) -------------------------------------
        wsU = wb.create_sheet(f"{p} Unitario")
        wsU.append([f"Economía unitaria — {idea.nombre} (precio de lista)"]); wsU["A1"].font = negrita
        wsU.append(["Concepto", "S/", "Fórmula"])
        u = {}

        def ufila(clave, concepto, formula):
            wsU.append([concepto, formula, formula]); u[clave] = f"'{p} Unitario'!$B${wsU.max_row}"

        ufila("precio", "Precio al cliente", f"={ref['precio']}")
        ufila("ingreso_neto", "Ingreso neto (sin IGV)", f"=IF({ref['igv_aplica']}=1,{u['precio']}/(1+{ref['tasa_igv']}),{u['precio']})")
        ufila("comision", "Comisión pasarela", f"={u['precio']}*{ref['tasa_pasarela']}+{ref['fijo_pasarela']}")
        ufila("impuesto", "Impuesto a la renta", f"={u['ingreso_neto']}*{ref['tasa_renta']}")
        ufila("cv_directo", "Costo variable directo", f"={ref['cv_directo']}")
        ufila("cv_total", "Costo variable total", f"={u['cv_directo']}+{u['comision']}+{u['impuesto']}")
        ufila("margen", "Margen por unidad", f"={u['ingreso_neto']}-{u['cv_total']}")
        ufila("margen_pct", "Margen % sobre ingreso neto", f"=IF({u['ingreso_neto']}=0,0,{u['margen']}/{u['ingreso_neto']})")
        ufila("fijos", "Costos fijos mensuales", f"={ref['fijos']}")
        ufila("pe_u", "Punto de equilibrio (unidades/mes)", f"=IF({u['margen']}>0,{u['fijos']}/{u['margen']},\"no alcanzable\")")
        ufila("pe_s", "Punto de equilibrio (S/ ventas/mes)", f"=IF(ISNUMBER({u['pe_u']}),{u['pe_u']}*{u['precio']},\"no alcanzable\")")
        ufila("inversion", "Inversión inicial", f"={ref['inversion']}")
        ufila("f1_inv", "F1: inversión ≤ tope", f"=IF({u['inversion']}<={G['tope_inv']},\"sí\",\"NO\")")
        ufila("f1_fijos", "F1: fijos ≤ tope mensual", f"=IF({u['fijos']}<={G['tope_fijos']},\"sí\",\"NO\")")
        ufila("u_sueldo1", "Unidades/mes para cubrir 1 sueldo", f"=IF({u['margen']}>0,({u['fijos']}+{G['sueldo']})/{u['margen']},\"no alcanzable\")")
        ufila("h_sueldo1", "Horas/semana para cubrir 1 sueldo", f"=IF(ISNUMBER({u['u_sueldo1']}),{ref['h_fijas']}+{ref['h_unidad']}*{u['u_sueldo1']}/({SEMANAS_POR_MES}),\"n/a\")")
        ufila("u_sueldo2", "Unidades/mes para cubrir 2 sueldos", f"=IF({u['margen']}>0,({u['fijos']}+2*{G['sueldo']})/{u['margen']},\"no alcanzable\")")
        ufila("h_sueldo2", "Horas/semana para cubrir 2 sueldos", f"=IF(ISNUMBER({u['u_sueldo2']}),{ref['h_fijas']}+{ref['h_unidad']}*{u['u_sueldo2']}/({SEMANAS_POR_MES}),\"n/a\")")
        # variantes de régimen y pasarela
        wsU.append([]); wsU.append(["Variantes (mismo precio de lista)", "Margen/u", "Fórmula"]); wsU[f"A{wsU.max_row}"].font = negrita
        ufila("v_rus", "Nuevo RUS cat. 1: sin IGV, sin renta %, +cuota como fijo → margen",
              f"={u['precio']}-({u['cv_directo']}+{u['comision']})")
        ufila("v_rus_pe", "Nuevo RUS: punto de equilibrio (u/mes)", f"=IF({u['v_rus']}>0,({u['fijos']}+{G['rus']})/{u['v_rus']},\"no alcanzable\")")
        ufila("v_4ta", "4.ª categoría: sin IGV, retención 8 % → margen", f"={u['precio']}-({u['cv_directo']}+{u['comision']}+{u['precio']}*{G['t4']})")
        ufila("v_culqi", "RER + pasarela Culqi → margen",
              f"={u['ingreso_neto']}-({u['cv_directo']}+{u['precio']}*{G['pas_t']}+{G['pas_f']}+{u['impuesto']})")
        if idea.rango_precio:
            for k_, lab in (("precio_min", "mínimo"), ("precio_max", "máximo")):
                pr = ref[k_]
                ing = f"IF({ref['igv_aplica']}=1,{pr}/(1+{ref['tasa_igv']}),{pr})"
                ufila(f"rp_{k_}", f"Margen/u con precio {lab} del rango",
                      f"={ing}-({u['cv_directo']}+{pr}*{ref['tasa_pasarela']}+{ref['fijo_pasarela']}+{ing}*{ref['tasa_renta']})")
        wsU.column_dimensions["A"].width = 60; wsU.column_dimensions["B"].width = 16; wsU.column_dimensions["C"].width = 90
        for r_ in range(3, wsU.max_row + 1):
            v = wsU[f"B{r_}"].value
            wsU[f"C{r_}"].value = ("'" + str(v)) if isinstance(v, str) and v.startswith("=") else None

        # --- Flujo 12 meses ------------------------------------------------
        wsF = wb.create_sheet(f"{p} Flujo")
        wsF.append([f"Flujo de caja a 12 meses — {idea.nombre}. Escenarios, no predicciones."]); wsF["A1"].font = negrita
        flujo_refs = {}
        for nm in ("pesimista", "base", "optimista"):
            wsF.append([]); wsF.append([f"Escenario {nm}"]); wsF[f"A{wsF.max_row}"].font = negrita
            r0 = wsF.max_row + 1
            wsF.append(["Precio escenario", f"={ref['precio']}*{ref[nm + '.f']}"])
            wsF.append(["Ingreso neto unitario", f"=IF({ref['igv_aplica']}=1,B{r0}/(1+{ref['tasa_igv']}),B{r0})"])
            wsF.append(["Costo variable total unitario", f"={ref['cv_directo']}+B{r0}*{ref['tasa_pasarela']}+{ref['fijo_pasarela']}+B{r0+1}*{ref['tasa_renta']}"])
            wsF.append(["Costos fijos mensuales", f"={ref['fijos']}"])
            wsF.append(["Inversión inicial", f"={ref['inversion']}"])
            rp, rin, rcv, rfj, rinv = r0, r0 + 1, r0 + 2, r0 + 3, r0 + 4
            cab = wsF.max_row + 1
            wsF.append(["Mes"] + list(range(1, MESES + 1)) + ["Total 12 m"])
            for c in range(1, MESES + 3):
                wsF.cell(row=cab, column=c).font = negrita; wsF.cell(row=cab, column=c).fill = gris
            ru, ri, rv, rf, rn, ra, rr, rh, rc, rs = (cab + i for i in range(1, 11))
            for lab in ("Unidades", "Ingresos netos", "Costos variables", "Costos fijos", "Flujo neto",
                        "Acumulado (tras inversión)", "Mes en que acumulado ≥ 0", "Horas/semana",
                        "Capital disponible (inicial + aportes)", "Mes en que flujo neto ≥ 1 sueldo"):
                wsF.append([lab])
            for m in range(1, MESES + 1):
                col = get_column_letter(m + 1)
                prev = get_column_letter(m)
                wsF[f"{col}{ru}"] = f"={ref[nm + '.u1']}*(1+{ref[nm + '.g']})^({m}-1)"
                wsF[f"{col}{ri}"] = f"={col}{ru}*$B${rin}"
                wsF[f"{col}{rv}"] = f"={col}{ru}*$B${rcv}"
                wsF[f"{col}{rf}"] = f"=$B${rfj}"
                wsF[f"{col}{rn}"] = f"={col}{ri}-{col}{rv}-{col}{rf}"
                wsF[f"{col}{ra}"] = f"=-$B${rinv}+{col}{rn}" if m == 1 else f"={prev}{ra}+{col}{rn}"
                wsF[f"{col}{rr}"] = f"=IF({col}{ra}>=0,{m},\"\")"
                wsF[f"{col}{rh}"] = f"={ref['h_fijas']}+{ref['h_unidad']}*{col}{ru}/({SEMANAS_POR_MES})"
                wsF[f"{col}{rc}"] = f"={G['capital']}+{G['aporte']}*{m}"
                wsF[f"{col}{rs}"] = f"=IF({col}{rn}>={G['sueldo']},{m},\"\")"
            tot = get_column_letter(MESES + 2)
            last = get_column_letter(MESES + 1)
            for r_ in (ru, ri, rv, rf, rn):
                wsF[f"{tot}{r_}"] = f"=SUM(B{r_}:{last}{r_})"
            wsF[f"{tot}{ra}"] = f"={last}{ra}"
            wsF[f"{tot}{rr}"] = f"=IF(COUNT(B{rr}:{last}{rr})=0,\"no en 12 meses\",MIN(B{rr}:{last}{rr}))"
            wsF[f"{tot}{rh}"] = f"=MAX(B{rh}:{last}{rh})"
            wsF[f"{tot}{rc}"] = f"=IF(SUMPRODUCT(--(-B{ra}:{last}{ra}>B{rc}:{last}{rc}))=0,\"cabe en el capital\",\"NO cabe\")"
            wsF[f"{tot}{rs}"] = f"=IF(COUNT(B{rs}:{last}{rs})=0,\"no en 12 meses\",MIN(B{rs}:{last}{rs}))"
            for r_ in (ra, rr, rh, rc, rs):
                wsF[f"{tot}{r_}"].font = negrita
            flujo_refs[nm] = {"resultado": f"'{p} Flujo'!${tot}${ra}", "recupera": f"'{p} Flujo'!${tot}${rr}",
                              "unidades": f"'{p} Flujo'!${tot}${ru}", "sueldo": f"'{p} Flujo'!${tot}${rs}",
                              "cabe": f"'{p} Flujo'!${tot}${rc}"}
        wsF.column_dimensions["A"].width = 38
        for c in range(2, MESES + 3):
            wsF.column_dimensions[get_column_letter(c)].width = 12

        # --- Sensibilidad (fórmula cerrada sobre el escenario base) --------
        wsX = wb.create_sheet(f"{p} Sensibilidad")
        wsX.append([f"Sensibilidad — {idea.nombre} — escenario base, ±20 % una variable a la vez"]); wsX["A1"].font = negrita
        wsX.append(["Resultado 12 m = Σ unidades × margen unitario − 12 × fijos − inversión; Σ unidades = u1·((1+g)^12−1)/g"])
        wsX.append(["Variación", 0.2]); wsX["B3"].fill = amarillo
        var = "$B$3"
        wsX.append(["Variable", "Valor base", "Resultado base", "Resultado −var", "Resultado +var", "Rango (abs)"])
        for c in "ABCDEF":
            wsX[f"{c}4"].font = negrita; wsX[f"{c}4"].fill = gris

        def R(precio=None, cv=None, fijos=None, inv=None, u1=None, g=None, tp=None, tr=None) -> str:
            precio = precio or f"({ref['precio']}*{ref['base.f']})"
            cv = cv or ref["cv_directo"]; fijos = fijos or ref["fijos"]; inv = inv or ref["inversion"]
            u1 = u1 or ref["base.u1"]; g = g or ref["base.g"]
            tp = tp or ref["tasa_pasarela"]; tr = tr or ref["tasa_renta"]
            ingreso = f"IF({ref['igv_aplica']}=1,{precio}/(1+{ref['tasa_igv']}),{precio})"
            margen = f"({ingreso}-({cv}+{precio}*{tp}+{ref['fijo_pasarela']}+{ingreso}*{tr}))"
            unidades = f"IF({g}=0,{MESES}*{u1},{u1}*((1+{g})^{MESES}-1)/{g})"
            return f"{unidades}*{margen}-{MESES}*{fijos}-{inv}"

        variables = [
            ("precio", f"{ref['precio']}", "precio"),
            ("unidades_mes1", ref["base.u1"], "u1"),
            ("crecimiento", ref["base.g"], "g"),
            ("cv_directo", ref["cv_directo"], "cv"),
            ("fijos", ref["fijos"], "fijos"),
            ("inversion", ref["inversion"], "inv"),
            ("tasa_pasarela", ref["tasa_pasarela"], "tp"),
            ("tasa_impuesto_renta", ref["tasa_renta"], "tr"),
        ]
        for nombre, celda, arg in variables:
            menos = f"({celda}*(1-{var}))"; mas = f"({celda}*(1+{var}))"
            if arg == "precio":
                menos = f"({menos}*{ref['base.f']})"; mas = f"({mas}*{ref['base.f']})"
            wsX.append([nombre, f"={celda}", f"={R()}", f"={R(**{arg: menos})}", f"={R(**{arg: mas})}", None])
            r_ = wsX.max_row
            wsX[f"F{r_}"] = f"=ABS(E{r_}-D{r_})"
        wsX.append([]); wsX.append(["La variable con mayor 'Rango' es la que más mueve el resultado."])
        for col, w in (("A", 22), ("B", 14), ("C", 16), ("D", 16), ("E", 16), ("F", 14)):
            wsX.column_dimensions[col].width = w

        comp_refs.append((idea, ref, u, flujo_refs))

    # --- Comparación ---------------------------------------------------------
    wsC = wb.create_sheet("Comparacion", 1)
    wsC.append(["Comparación con el mismo método — escenario base salvo donde se indica"]); wsC["A1"].font = negrita
    cab = ["Idea", "Inversión inicial", "Margen por unidad", "Margen %", "Punto de equilibrio (u/mes)",
           "Recupera (mes, base)", "Resultado 12 m (pesimista)", "Resultado 12 m (base)", "Resultado 12 m (optimista)",
           "Horas/semana (pico base)", "F1 inversión", "F1 fijos", "Unid./mes 1 sueldo", "Horas/sem 1 sueldo",
           "Mes cubre 1 sueldo (base)", "Mes cubre 1 sueldo (optimista)", "Cabe en capital (base)"]
    wsC.append(cab)
    for c in range(1, len(cab) + 1):
        wsC.cell(row=2, column=c).font = negrita; wsC.cell(row=2, column=c).fill = gris
    for idea, ref, u, fr in comp_refs:
        wsC.append([idea.nombre, f"={u['inversion']}", f"={u['margen']}", f"={u['margen_pct']}", f"={u['pe_u']}",
                    f"={fr['base']['recupera']}", f"={fr['pesimista']['resultado']}", f"={fr['base']['resultado']}",
                    f"={fr['optimista']['resultado']}", f"={ref['horas']}", f"={u['f1_inv']}", f"={u['f1_fijos']}",
                    f"={u['u_sueldo1']}", f"={u['h_sueldo1']}", f"={fr['base']['sueldo']}", f"={fr['optimista']['sueldo']}",
                    f"={fr['base']['cabe']}"])
    wsC.column_dimensions["A"].width = 52
    for c in range(2, len(cab) + 1):
        wsC.column_dimensions[get_column_letter(c)].width = 20
    for row in wsC.iter_rows(min_row=3, max_row=wsC.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    wb.save(ruta)


# =============================================================================
# 6b. DESEMPATE A vs B (Fase 4) — cálculo e impresión
# =============================================================================


def _valores(datos: dict[str, dict[str, Dato]]) -> tuple[dict, dict]:
    """Convierte {A: {clave: Dato}, B: {...}} en dos dicts de valores planos."""
    return {k: d.valor for k, d in datos["A"].items()}, {k: d.valor for k, d in datos["B"].items()}


def desempate(A: dict, B: dict, *, cv_s29: float = CV_S29_CHATBOT.valor,
              horas_construccion_tope: float = HORAS_CONSTRUCCION_A3_TOPE.valor,
              horas_cliente_mes: float = HORAS_POR_CLIENTE_MES_A.valor,
              meses_recurrencia: float = MESES_RECURRENCIA_A.valor,
              umbral_empate: float = UMBRAL_EMPATE.valor) -> dict:
    """Margen real por hora (MRH) de A y B con la fórmula exacta del plan de validación, sección D.

    A y B son dicts de valores planos con las claves de DATOS_DESEMPATE (None = falta).
    Devuelve un dict con `faltan` (lista de claves sin dato), `A`, `B` (desglose y MRH cuando se puede calcular)
    y `recomendacion` (texto de la regla del plan, o None si faltan datos).
    Nunca inventa: con cualquier dato necesario en None, deja el MRH en None y lista qué falta.
    """
    faltan: list[str] = []

    def req(d: dict, lado: str, clave: str):
        v = d.get(clave)
        if v is None:
            faltan.append(f"{lado}.{clave}")
        return v

    # ---- A -----------------------------------------------------------------
    paso_A = req(A, "A", "paso_A2")
    adel_A = req(A, "A", "adelantos_cobrados")
    cost_A = req(A, "A", "costos_directos_piloto")
    mens_A = req(A, "A", "mensualidad_acordada")
    cli_A = req(A, "A", "clientes_siguen_mes2")
    h_val_A = req(A, "A", "horas_A0_A1_A2")
    a3_ok = req(A, "A", "a3_terminado")
    h_con_A = A.get("horas_construccion_A3")
    h_con_origen = None
    if a3_ok is True:
        if h_con_A is None:
            faltan.append("A.horas_construccion_A3 (A3 terminó: hacen falta las horas cronometradas)")
        else:
            h_con_origen = "cronometradas en A3"
    elif a3_ok is False:
        h_con_A = horas_construccion_tope
        h_con_origen = f"{horas_construccion_tope:.0f} h por regla del plan (A3 no terminó)"

    resA: dict = {"paso": paso_A, "mrh": None}
    if None not in (adel_A, cost_A, mens_A, cli_A, h_val_A, h_con_A):
        margen_rec = meses_recurrencia * (mens_A - cv_s29) * cli_A
        num = adel_A - cost_A + margen_rec
        h_cli = horas_cliente_mes * cli_A * meses_recurrencia
        den = h_val_A + h_con_A + h_cli
        resA.update({
            "adelantos": adel_A, "costos_directos": cost_A, "mensualidad": mens_A, "cv_s29": cv_s29,
            "clientes_siguen": cli_A, "margen_recurrente_2m": margen_rec, "numerador": num,
            "horas_A0_A1_A2": h_val_A, "horas_construccion": h_con_A, "horas_construccion_origen": h_con_origen,
            "horas_clientes": h_cli, "denominador": den,
            "mrh": (num / den) if den > 0 else None,
            "mrh_sin_horas_validacion": (num / (h_con_A + h_cli)) if (h_con_A + h_cli) > 0 else None,
        })
        if den <= 0:
            faltan.append("A: el total de horas es 0; no se puede dividir")

    # ---- B -----------------------------------------------------------------
    paso_B = req(B, "B", "paso_B2")
    adel_B = req(B, "B", "adelantos_cobrados")
    cost_B = req(B, "B", "costos_reales_pedidos")
    h_B = req(B, "B", "horas_B1_B2_B3_B4")
    resB: dict = {"paso": paso_B, "mrh": None}
    if None not in (adel_B, cost_B, h_B):
        num = adel_B - cost_B
        resB.update({"adelantos": adel_B, "costos": cost_B, "numerador": num, "horas": h_B,
                     "mrh": (num / h_B) if h_B > 0 else None})
        if h_B <= 0:
            faltan.append("B: el total de horas es 0; no se puede dividir")

    # ---- Regla del plan ----------------------------------------------------
    rec = None
    regla = None
    if paso_A is not None and paso_B is not None:
        if paso_A and not paso_B:
            rec, regla = "A (chatbot)", "regla 1: solo A pasó su criterio A2"
        elif paso_B and not paso_A:
            rec, regla = "B (detalles)", "regla 1: solo B pasó su criterio B2"
        elif not paso_A and not paso_B:
            rec, regla = "volver a Fase 3", "regla 3: ninguna pasó su criterio; no se fuerza la elección"
        elif resA["mrh"] is not None and resB["mrh"] is not None:
            ma, mb = resA["mrh"], resB["mrh"]
            mayor = max(ma, mb)
            dif = abs(ma - mb) / mayor if mayor > 0 else 0.0
            if mayor <= 0:
                rec, regla = "volver a Fase 3", "regla 2: ambas pasaron pero ninguna tiene MRH positivo; la agencia recomienda reabrir"
            elif dif < umbral_empate:
                rec, regla = "A (chatbot)", f"regla 2: ambas pasaron; diferencia {dif:.1%} < {umbral_empate:.0%} → gana A por recurrencia"
            elif ma > mb:
                rec, regla = "A (chatbot)", f"regla 2: ambas pasaron; MRH A > B, diferencia {dif:.1%}"
            else:
                rec, regla = "B (detalles)", f"regla 2: ambas pasaron; MRH B > A, diferencia {dif:.1%}"
            resA["diferencia_relativa"] = resB["diferencia_relativa"] = dif
    return {"faltan": faltan, "A": resA, "B": resB, "recomendacion": rec, "regla": regla}


def _si_no(v) -> str:
    return "—" if v is None else ("sí" if v else "no")


def imprimir_desempate(res: dict, datos: dict[str, dict[str, Dato]], titulo: str) -> None:
    print(f"\n=== DESEMPATE A vs B — {titulo} ===")
    print("Fórmula: entregables/plan-de-validacion.md, sección D. Comando: python3 finanzas/modelo.py --desempate [--ejemplo]")
    print("\nEntradas (valor · origen):")
    for lado in ("A", "B"):
        for k, d in datos[lado].items():
            v = "—" if d.valor is None else (_si_no(d.valor) if isinstance(d.valor, bool) else f"{d.valor:g}")
            print(f"  {lado}.{k:<24} {v:>8}  [{d.origen}] {d.descripcion}")
    print(f"  constante CV S-29        {CV_S29_CHATBOT.valor:>8.2f}  [{CV_S29_CHATBOT.origen}]")
    if res["faltan"]:
        print("\nFALTAN DATOS — no se calcula nada con valores inventados. Registrar en DATOS_DESEMPATE o pasar con --dato:")
        for f in res["faltan"]:
            print(f"  - {f}")
    a, b = res["A"], res["B"]
    if a.get("mrh") is not None:
        print("\nA (chatbot):")
        print(f"  numerador   = {s(a['adelantos'])} − {s(a['costos_directos'])} + {MESES_RECURRENCIA_A.valor:.0f} × ({s(a['mensualidad'])} − {s(a['cv_s29'])}) × {a['clientes_siguen']:g}"
              f" = {s(a['numerador'])}")
        print(f"  denominador = {n(a['horas_A0_A1_A2'])} h (A0+A1+A2) + {n(a['horas_construccion'])} h ({a['horas_construccion_origen']})"
              f" + {HORAS_POR_CLIENTE_MES_A.valor:.0f} h × {a['clientes_siguen']:g} × {MESES_RECURRENCIA_A.valor:.0f} m = {n(a['denominador'])} h")
        print(f"  MRH A = {a['mrh']:.2f} S/ por hora   (pasó A2: {_si_no(a['paso'])})")
    if b.get("mrh") is not None:
        print("\nB (detalles):")
        print(f"  numerador   = {s(b['adelantos'])} − {s(b['costos'])} = {s(b['numerador'])}")
        print(f"  denominador = {n(b['horas'])} h (B1+B2+B3+B4)")
        print(f"  MRH B = {b['mrh']:.2f} S/ por hora   (pasó B2: {_si_no(b['paso'])})")
    if res["recomendacion"]:
        print(f"\nRECOMENDACIÓN DE LA AGENCIA: {res['recomendacion']}  ({res['regla']})")
        print("La decisión es de los socios y se registra en expediente/decisiones.md.")
    elif not res["faltan"]:
        print("\nSin recomendación: faltan los MRH para aplicar la regla 2.")


def verificar_ilustracion_plan() -> None:
    """Corre la fórmula con las cifras ILUSTRATIVAS del plan y las contrasta con lo que el plan dice (≈ 6, ≈ 10, ≈ 10)."""
    A, B = _valores(DATOS_DESEMPATE_EJEMPLO)
    r1 = desempate(A, B)
    A2 = dict(A, adelantos_cobrados=2 * A["adelantos_cobrados"], costos_directos_piloto=2 * A["costos_directos_piloto"],
              clientes_siguen_mes2=2)
    r2 = desempate(A2, B)
    filas = [
        ("A con 1 piloto", r1["A"]["mrh"], r1["A"]["mrh_sin_horas_validacion"], MRH_PLAN_ILUSTRATIVO["A_1_piloto"]),
        ("A con 2 pilotos", r2["A"]["mrh"], r2["A"]["mrh_sin_horas_validacion"], MRH_PLAN_ILUSTRATIVO["A_2_pilotos"]),
        ("B con 6 pedidos", r1["B"]["mrh"], r1["B"]["mrh"], MRH_PLAN_ILUSTRATIVO["B"]),
    ]
    print("\n--- Verificación contra la ilustración del plan (sección D) ---")
    print(f"  {'caso':<18} {'fórmula del plan':>17} {'sin h A0–A2':>12} {'plan dice ≈':>12} {'coincide':>9}")
    tol = 1.0  # S/ por hora: "aproximadamente"
    for nombre, mrh, mrh_alt, esperado in filas:
        ok = abs(mrh - esperado) <= tol
        print(f"  {nombre:<18} {mrh:>17.2f} {mrh_alt:>12.2f} {esperado:>12.1f} {('sí' if ok else 'NO'):>9}")
    print(f"  Tolerancia: ± {tol:.0f} S/ por hora. 'sin h A0–A2' = mismo numerador dividido solo entre construcción + horas por cliente.")
    desvios = [f for f in filas if abs(f[1] - f[3]) > tol]
    if desvios:
        print("  DIFERENCIA: la cifra ilustrativa del plan NO se reproduce con la fórmula del plan para: "
              + ", ".join(f[0] for f in desvios) + ".")
        print("  La fórmula escrita en la sección D manda (incluye las horas de A0+A1+A2 en el denominador); la cifra ilustrativa")
        print("  del plan corresponde a omitir esas horas (columna 'sin h A0–A2'). El Director debe corregir la ilustración en el plan,")
        print("  no la fórmula. Este script usa siempre la fórmula.")
    else:
        print("  Las cifras ilustrativas del plan se reproducen con la fórmula.")
    print("  Con la fórmula del plan y estas cifras ILUSTRATIVAS, A con 2 pilotos da "
          f"{r2['A']['mrh']:.2f} S/ por hora frente a {r1['B']['mrh']:.2f} de B; el desempate no está tan abierto como dice el plan.")


# =============================================================================
# 7. CLI
# =============================================================================


def _parse_valor(texto: str):
    t = texto.strip().lower()
    if t in ("si", "sí", "true", "1", "yes"):
        return True
    if t in ("no", "false", "0"):
        return False
    return float(texto.replace(",", "."))


def main_desempate(args) -> int:
    if args.ejemplo:
        datos = DATOS_DESEMPATE_EJEMPLO
        titulo = "cifras ILUSTRATIVAS del plan (NO son datos reales)"
    else:
        datos = {lado: dict(v) for lado, v in DATOS_DESEMPATE.items()}
        titulo = "datos reales de los socios (DATOS_DESEMPATE / --dato)"
    for item in args.dato:
        try:
            clave, valor = item.split("=", 1)
            lado, nombre = clave.strip().split(".", 1)
            if nombre not in datos[lado]:
                raise KeyError(nombre)
            d = datos[lado][nombre]
            datos[lado][nombre] = Dato(_parse_valor(valor), "VERIFICADO: registro de los socios pasado por --dato " + "(" + item + ")", d.descripcion)
        except (ValueError, KeyError) as exc:
            claves = ", ".join(f"{l}.{k}" for l in datos for k in datos[l])
            print(f"--dato inválido '{item}' ({exc}). Claves válidas: {claves}", file=sys.stderr)
            return 2
    cv_modelo = sum(e.valor for e in IDEAS[0].costo_variable_unitario.values()) if IDEAS else None
    print(banner(), end="")
    A, B = _valores(datos)
    res = desempate(A, B)
    imprimir_desempate(res, datos, titulo)
    if cv_modelo is not None and abs(cv_modelo - CV_S29_CHATBOT.valor) > 0.01:
        print(f"\nAVISO: el costo variable de la idea 1 en el modelo ({cv_modelo:.2f}) ya no coincide con la constante S-29 del plan "
              f"({CV_S29_CHATBOT.valor:.2f}); el plan manda, pero conviene revisar.")
    if args.ejemplo:
        verificar_ilustracion_plan()
    print(banner(), end="")
    return 0 if not res["faltan"] else 3


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--comparar", action="store_true", help="solo la tabla comparativa")
    ap.add_argument("--xlsx", nargs="?", const="finanzas/modelo.xlsx", metavar="RUTA",
                    help="regenerar la hoja de cálculo con fórmulas (por defecto finanzas/modelo.xlsx)")
    ap.add_argument("--json", action="store_true", help="resultados en JSON")
    ap.add_argument("--idea", help="evaluar solo la idea con este nombre (coincidencia parcial)")
    ap.add_argument("--desempate", action="store_true",
                    help="Fase 4: margen real por hora de A (chatbot) y B (detalles) con DATOS_DESEMPATE")
    ap.add_argument("--ejemplo", action="store_true",
                    help="con --desempate: usar las cifras ILUSTRATIVAS del plan y verificarlas (no son datos)")
    ap.add_argument("--dato", action="append", default=[], metavar="LADO.CLAVE=VALOR",
                    help="con --desempate: cargar un dato real por argumento, ej. A.adelantos_cobrados=100 o B.paso_B2=si")
    args = ap.parse_args(argv)

    if args.ejemplo and not args.desempate:
        print("--ejemplo solo tiene sentido con --desempate.", file=sys.stderr)
        return 2
    if args.desempate:
        return main_desempate(args)

    ideas = IDEAS
    if args.idea:
        ideas = [i for i in IDEAS if args.idea.lower() in i.nombre.lower()]
        if not ideas:
            print(f"No hay ninguna idea cuyo nombre contenga '{args.idea}'.", file=sys.stderr)
            return 2
    if not ideas:
        print("IDEAS está vacío: cargue las ideas reales con su origen.", file=sys.stderr)
        return 2

    errores = validar(ideas)
    if errores:
        print("ERRORES DE VALIDACIÓN — el modelo no corre hasta corregirlos:", file=sys.stderr)
        for e in errores:
            print("  - " + e, file=sys.stderr)
        return 1

    resultados = [evaluar(i) for i in ideas]

    if args.json:
        print(json.dumps({"es_ejemplo": ES_EJEMPLO, "ideas": resultados}, ensure_ascii=False, indent=2,
                         default=lambda o: None if (isinstance(o, float) and not math.isfinite(o)) else str(o)))
        return 0

    print(banner(), end="")
    imprimir_globales()
    if args.comparar:
        imprimir_comparacion(ideas, resultados)
    else:
        for idea, r in zip(ideas, resultados):
            imprimir_idea(idea, r)
        if len(ideas) > 1:
            imprimir_comparacion(ideas, resultados)
    print(banner(), end="")

    if args.xlsx:
        exportar_xlsx(ideas, args.xlsx)
        print(f"\nHoja de cálculo regenerada: {args.xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
